from __future__ import annotations

import json
import shutil
from copy import deepcopy
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree


REPO_ROOT = Path(__file__).resolve().parents[1]
WORK_ROOT = REPO_ROOT.parent
TEMPLATE_DOC = Path(r"C:\Users\laksh\Downloads\Abstract.docx")
RESULTS_DIR = REPO_ROOT / "results"
REPORT_DIR = REPO_ROOT / "report"
ASSET_DIR = REPORT_DIR / "assets"
OUTPUTS_DIR = WORK_ROOT / "outputs" / "athlete_final_report_final"
OUT_DOCX = REPORT_DIR / "Athlete_Injury_Risk_Prediction_Final_Report.docx"
OUTPUT_COPY = OUTPUTS_DIR / "Athlete_Injury_Risk_Prediction_Final_Report.docx"
DOWNLOAD_COPY = Path(r"C:\Users\laksh\Downloads\Athlete_Injury_Risk_Prediction_Final_Report.docx")


def ensure_dirs() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)


def load_results() -> tuple[dict, pd.DataFrame, pd.DataFrame]:
    summary = json.loads((RESULTS_DIR / "experiment_summary.json").read_text(encoding="utf-8"))
    holdout = pd.read_csv(RESULTS_DIR / "holdout_leaderboard.csv")
    cv = pd.read_csv(RESULTS_DIR / "cross_validation_summary.csv")
    return summary, holdout, cv


def copy_columns(source_section, target_section) -> None:
    source_cols = source_section._sectPr.xpath("./w:cols")
    if not source_cols:
        return
    source_cols = source_cols[0]
    sect_pr = target_section._sectPr
    existing = sect_pr.xpath("./w:cols")
    for node in existing:
        sect_pr.remove(node)
    sect_pr.append(deepcopy(source_cols))


def apply_section_layout(source_section, target_section, copy_col_spec: bool) -> None:
    target_section.page_width = source_section.page_width
    target_section.page_height = source_section.page_height
    target_section.left_margin = source_section.left_margin
    target_section.right_margin = source_section.right_margin
    target_section.top_margin = source_section.top_margin
    target_section.bottom_margin = source_section.bottom_margin
    target_section.header_distance = source_section.header_distance
    target_section.footer_distance = source_section.footer_distance
    if copy_col_spec:
        copy_columns(source_section, target_section)


def configure_document(doc: Document) -> None:
    source = Document(TEMPLATE_DOC)
    apply_section_layout(source.sections[0], doc.sections[0], copy_col_spec=True)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)


def start_two_column_body(doc: Document) -> None:
    source = Document(TEMPLATE_DOC)
    section_two = doc.add_section(WD_SECTION.CONTINUOUS)
    apply_section_layout(source.sections[1], section_two, copy_col_spec=True)
    break_paragraph = doc.paragraphs[-1]
    break_paragraph.paragraph_format.space_before = Pt(0)
    break_paragraph.paragraph_format.space_after = Pt(0)


def shade_cell(cell, fill: str = "EAEAEA") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("left", "top", "right", "bottom"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_cell_margins(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table) -> None:
    border = {"val": "single", "sz": 6, "space": 0, "color": "808080"}
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)
            set_cell_margins(cell)
            if row_idx == 0:
                shade_cell(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_column_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_title(doc: Document, text: str) -> None:
    paragraph = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(21)
    run.bold = True


def add_author_block(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def add_abstract_block(doc: Document, abstract_text: str, index_terms: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    bold = paragraph.add_run("Abstract-")
    bold.bold = True
    bold.font.name = "Times New Roman"
    bold.font.size = Pt(9)
    run = paragraph.add_run(abstract_text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)

    idx = doc.add_paragraph()
    idx.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    idx.paragraph_format.space_after = Pt(6)
    idx_bold = idx.add_run("Index Terms-")
    idx_bold.bold = True
    idx_bold.font.name = "Times New Roman"
    idx_bold.font.size = Pt(9)
    idx_run = idx.add_run(index_terms)
    idx_run.font.name = "Times New Roman"
    idx_run.font.size = Pt(9)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.bold = True
    run.small_caps = True


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.16)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def add_caption(doc: Document, text: str, above: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2 if above else 0)
    paragraph.paragraph_format.space_after = Pt(3 if above else 6)
    paragraph.paragraph_format.keep_with_next = above
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)
    if above:
        run.bold = True
    else:
        run.italic = True


def add_reference(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(8.5)


def add_figure(doc: Document, image_path: Path, width_inches: float, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    add_caption(doc, caption, above=False)


def build_tradeoff_chart(holdout_df: pd.DataFrame) -> Path:
    order = [
        "baseline_svm_rbf_rfe10",
        "random_forest_balanced",
        "logistic_extension_all_features",
        "baseline_svm_threshold_rule_0_085",
    ]
    labels = ["SVM", "RF", "Logistic", "Threshold"]
    frame = holdout_df.set_index("model").loc[order].reset_index()
    accuracy = frame["accuracy"].tolist()
    balanced = frame["balanced_accuracy"].tolist()
    recall = frame["recall"].tolist()

    out = ASSET_DIR / "holdout_tradeoff_chart.png"
    x = range(len(labels))
    width = 0.23
    fig, ax = plt.subplots(figsize=(5.6, 3.35), dpi=180)
    ax.bar([i - width for i in x], accuracy, width=width, label="Accuracy", color="#1F4E79")
    ax.bar(x, balanced, width=width, label="Balanced Accuracy", color="#C0504D")
    ax.bar([i + width for i in x], recall, width=width, label="Recall", color="#9BBB59")
    ax.set_xticks(list(x), labels=labels)
    ax.set_ylim(0, 1.08)
    ax.set_ylabel("Score")
    ax.set_title("Holdout Metric Tradeoff", fontsize=10)
    ax.grid(axis="y", linestyle="--", linewidth=0.5, alpha=0.4)
    ax.legend(fontsize=7, frameon=False, ncols=3, loc="upper left")
    for bars in ax.containers:
        ax.bar_label(bars, fmt="%.2f", fontsize=7, padding=2)
    plt.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)
    return out


def build_dataset_table(doc: Document) -> None:
    add_caption(doc, "Table I. Dataset summary and split structure used in the final experiment.", above=True)
    rows = [
        ("Item", "Value"),
        ("Total records", "200"),
        ("Original columns", "17"),
        ("Target variable", "Injury_Indicator"),
        ("Injury prevalence", "14 of 200 records (7%)"),
        ("Missing values", "None"),
        ("Baseline modeling fields", "13 numeric features after dropping Athlete_ID, Gender, and Position"),
        ("All-feature extension", "15 predictors after dropping Athlete_ID only"),
        ("Holdout split", "140 train / 30 validation / 30 test"),
        ("Class counts by split", "Train 130/10, validation 28/2, test 28/2 (non-injury/injury)"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, (left, right) in enumerate(rows):
        left_cell = table.rows[row_idx].cells[0]
        right_cell = table.rows[row_idx].cells[1]
        left_p = left_cell.paragraphs[0]
        right_p = right_cell.paragraphs[0]
        left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        right_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        left_run = left_p.add_run(left)
        right_run = right_p.add_run(right)
        left_run.font.name = "Times New Roman"
        right_run.font.name = "Times New Roman"
        left_run.font.size = Pt(8)
        right_run.font.size = Pt(8)
        if row_idx == 0:
            left_run.bold = True
            right_run.bold = True
    set_repeat_table_header(table.rows[0])
    set_column_widths(table, [1.15, 2.2])
    style_table(table)


def build_selected_features_table(doc: Document) -> None:
    rows = [
        ("Feature", "Reason it was useful in the baseline"),
        ("Height (cm)", "Body profile variable that can interact with movement and exposure"),
        ("Training intensity", "Direct workload signal"),
        ("Training hours/week", "Total training volume"),
        ("Recovery days/week", "Recovery opportunity across the week"),
        ("Rest between events", "Short-term rest window"),
        ("Fatigue score", "Accumulated fatigue signal"),
        ("Performance score", "Current output or form"),
        ("Team contribution", "Proxy for playing role and match involvement"),
        ("Load balance", "How evenly workload is distributed"),
        ("ACL risk score", "Direct injury-risk indicator"),
    ]
    add_caption(doc, "Table II. Features selected by RFE for the baseline SVM pipeline.", above=True)
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, (left, right) in enumerate(rows):
        left_cell = table.rows[row_idx].cells[0]
        right_cell = table.rows[row_idx].cells[1]
        left_p = left_cell.paragraphs[0]
        right_p = right_cell.paragraphs[0]
        left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        right_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        left_run = left_p.add_run(left)
        right_run = right_p.add_run(right)
        left_run.font.name = "Times New Roman"
        right_run.font.name = "Times New Roman"
        left_run.font.size = Pt(7.6)
        right_run.font.size = Pt(7.6)
        if row_idx == 0:
            left_run.bold = True
            right_run.bold = True
    set_repeat_table_header(table.rows[0])
    set_column_widths(table, [1.18, 2.17])
    style_table(table)


def build_holdout_table(doc: Document, holdout_df: pd.DataFrame) -> None:
    order = [
        "baseline_svm_rbf_rfe10",
        "random_forest_balanced",
        "logistic_extension_all_features",
        "baseline_svm_threshold_rule_0_085",
    ]
    labels = {
        "baseline_svm_rbf_rfe10": "Baseline SVM",
        "random_forest_balanced": "Random Forest",
        "logistic_extension_all_features": "Logistic extension",
        "baseline_svm_threshold_rule_0_085": "SVM threshold",
    }
    frame = holdout_df.set_index("model").loc[order].reset_index()
    add_caption(doc, "Table III. Holdout test-set comparison across the final decision rules.", above=True)
    table = doc.add_table(rows=len(frame) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Model", "Acc", "Prec", "Rec"]
    for idx, header in enumerate(headers):
        paragraph = table.rows[0].cells[idx].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.font.name = "Times New Roman"
        run.font.size = Pt(7.0)
        run.bold = True
    for row_idx, (_, row) in enumerate(frame.iterrows(), start=1):
        values = [
            labels[row["model"]],
            f"{row['accuracy']:.2f}",
            f"{row['precision']:.2f}",
            f"{row['recall']:.2f}",
        ]
        for col_idx, value in enumerate(values):
            paragraph = table.rows[row_idx].cells[col_idx].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            run.font.name = "Times New Roman"
            run.font.size = Pt(7.0)
    set_repeat_table_header(table.rows[0])
    set_column_widths(table, [1.42, 0.44, 0.44, 0.44])
    style_table(table)


def build_cv_table(doc: Document, cv_df: pd.DataFrame) -> None:
    order = [
        "random_forest_balanced",
        "logistic_extension_all_features",
        "logistic_extension_rfe8",
        "baseline_svm_rbf_rfe10",
    ]
    labels = {
        "random_forest_balanced": "Random Forest",
        "logistic_extension_all_features": "Logistic extension",
        "logistic_extension_rfe8": "Compact logistic",
        "baseline_svm_rbf_rfe10": "Baseline SVM",
    }
    frame = cv_df.set_index("model").loc[order].reset_index()
    add_caption(doc, "Table IV. Mean scores from repeated 5x10 stratified cross-validation.", above=True)
    table = doc.add_table(rows=len(frame) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Model", "Acc", "Bal", "Rec"]
    for idx, header in enumerate(headers):
        paragraph = table.rows[0].cells[idx].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.font.name = "Times New Roman"
        run.font.size = Pt(7.0)
        run.bold = True
    for row_idx, (_, row) in enumerate(frame.iterrows(), start=1):
        values = [
            labels[row["model"]],
            f"{row['accuracy']:.2f}",
            f"{row['balanced_accuracy']:.2f}",
            f"{row['recall']:.2f}",
        ]
        for col_idx, value in enumerate(values):
            paragraph = table.rows[row_idx].cells[col_idx].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            run.font.name = "Times New Roman"
            run.font.size = Pt(7.0)
    set_repeat_table_header(table.rows[0])
    set_column_widths(table, [1.52, 0.34, 0.48, 0.34])
    style_table(table)


def build_report() -> dict[str, str]:
    ensure_dirs()
    summary, holdout_df, cv_df = load_results()
    tradeoff_chart = build_tradeoff_chart(holdout_df)
    confusion_chart = RESULTS_DIR / "figures" / "baseline_svm_confusion_matrix.png"
    importance_chart = RESULTS_DIR / "figures" / "random_forest_feature_importance.png"

    doc = Document()
    configure_document(doc)

    add_title(doc, "Athlete Injury Risk Prediction: From an SVM Baseline to a Safer Early-Warning Pipeline")
    add_author_block(doc, "Lakshya Mohan Rastogi\nPennsylvania State University\nlxr5416@psu.edu")
    add_abstract_block(
        doc,
        "Athlete injury prediction looks like a standard binary classification problem, but in practice it is harder than that. Injury cases are rare, the cost of a missed case is high, and a model can look strong on accuracy while still failing the exact athletes it is supposed to protect. This project revisits an SVM-based injury prediction pipeline inspired by Li [1] and evaluates it on a classroom dataset of 200 collegiate athlete records containing 14 injury cases. The baseline workflow removes identifier and selected categorical fields, standardizes numeric variables, applies recursive feature elimination (RFE) to keep 10 predictors, and trains an RBF-kernel SVM. To move beyond a single-model result, the study also compares a balanced Random Forest, an all-feature balanced logistic regression pipeline, and a lower-threshold SVM warning rule. On the fixed holdout test split, the SVM and Random Forest each reached 96.67% accuracy but recovered only one of the two injury cases, giving recall of 0.50. The logistic extension reduced accuracy to 93.33% but detected both injury cases. The threshold rule also achieved recall of 1.00, but with three false alarms. Repeated 5x10 stratified cross-validation showed that Random Forest had the highest mean accuracy (0.9560), while the logistic extension produced the best overall balance between accuracy and injury recovery. The main conclusion is that injury prediction should not be judged by accuracy alone. Recall, false negatives, and model interpretability need to be treated as central outcomes.",
        "athlete injury prediction, machine learning, support vector machine, random forest, recall, sports analytics",
    )
    start_two_column_body(doc)

    add_heading(doc, "I. INTRODUCTION")
    add_body(
        doc,
        "Sports injuries affect much more than a single game. They change how an athlete trains, whether a team can keep a stable lineup, and how much confidence coaches have in their rotation. In many settings, an injury also carries financial consequences because time lost to recovery affects roster planning, player value, and overall performance. That is why injury prediction has become an important topic in sports analytics and sports medicine [2]-[4]."
    )
    add_body(
        doc,
        "Recent tracking systems have made this problem more measurable than before. Teams can now collect workload, recovery, movement, and performance information from training logs, match reports, and wearable devices. In theory, that makes machine learning a strong fit for the problem because the model can combine many weak signals instead of depending on one simple rule. In practice, however, injury prediction remains difficult because injury cases are rare, the causes are multi-factor, and the cost of a missed case is usually much higher than the cost of a false warning [3], [4]."
    )
    add_body(
        doc,
        "The parent paper for this project, written by Li [1], uses a support vector machine (SVM) together with big data framing and reports strong headline performance. That paper gave me a useful starting point because it treats injury prediction as a supervised classification task and combines feature selection with an RBF-kernel SVM. At the same time, it leaves open a practical question: if the same general pipeline is applied to a much smaller and more imbalanced classroom dataset, does high accuracy still mean the model is doing the job well?"
    )
    add_body(
        doc,
        "This final project answers that question by rebuilding the SVM-style pipeline, comparing it with a Random Forest and a balanced logistic regression extension, and then changing the evaluation lens. Instead of treating accuracy as the only bottom line, the project looks directly at recall, false negatives, balanced accuracy, and model interpretability. The result is a more honest view of what the current pipeline can do and where it still falls short."
    )

    add_heading(doc, "II. BACKGROUND AND RELATED WORK")
    add_body(
        doc,
        "A recent systematic review by Van Eetvelde et al. [2] shows that sport injury prediction already uses a mix of tree-based ensembles, SVMs, and neural networks, but performance varies widely from study to study. The review is useful because it makes two points that matter for this project. First, there is no single model family that wins in every setting. Second, many studies still struggle with data imbalance, interpretation, and generalizability. In other words, getting a high score on one dataset does not automatically mean the system is ready for practice."
    )
    add_body(
        doc,
        "Work in football and sports medicine supports the same idea. Majumdar et al. [3] argue that injury prediction has to be tied to training load, fatigue, and recovery rather than treated as a one-time label. Pareek et al. [4] make a similar point from a broader sports medicine perspective and note that clinician trust is a major barrier to adoption. A model that gives a label without a usable explanation is harder to act on, even if the raw metrics look strong."
    )
    add_body(
        doc,
        "This is also why time-aware methods are appealing. Ye et al. [5] propose a deep learning approach built around time-series image encoding, which lets the model learn how injury risk develops across sequences instead of isolated records. That direction is important for future work because many injury signals grow over days or weeks rather than appearing all at once. My current project does not have the longitudinal structure needed for that kind of model, but the paper still influenced the way I thought about future extensions."
    )
    add_body(
        doc,
        "One more issue from the literature directly affects how this project is evaluated. On imbalanced datasets, accuracy can make a classifier look better than it really is. Saito and Rehmsmeier [9] show why precision-recall based evaluation is often more informative than broader summary measures when positives are rare. That matters here because only 14 of the 200 athlete records belong to the injury class. A model can score well by mostly predicting non-injury, which is exactly why this report gives extra attention to recall and false negatives."
    )

    add_heading(doc, "III. RESEARCH QUESTION AND PROJECT FRAMING")
    add_body(
        doc,
        "The main research question for this project is straightforward: can an SVM-based injury prediction pipeline, inspired by the parent paper, still be considered useful once missed injuries are treated as a central outcome rather than a side metric? I was not trying to beat the parent paper on raw accuracy because the local dataset is much smaller and does not include the same big-data setting. The goal was to test whether the same basic modeling idea still holds up under a stricter and more practical evaluation."
    )
    add_body(
        doc,
        "The project makes four concrete changes relative to a simple baseline replication. First, it rebuilds the parent-paper style SVM pipeline on the local athlete dataset. Second, it adds a Random Forest comparison so the model behavior is not judged from one algorithm only. Third, it adds a balanced logistic extension that keeps the full feature space instead of dropping categorical variables. Fourth, it adds a recall-first threshold rule to show what happens when the system is used more like an early warning tool than a hard binary classifier. Those changes are important for the rubric because they move the work beyond a direct copy of the parent paper and create a clear contrast in both method and interpretation."
    )

    add_heading(doc, "IV. DATASET AND EXPERIMENTAL SETUP")
    add_body(
        doc,
        "The project uses a local classroom dataset named collegiate_athlete_injury_dataset.csv. It contains 200 athlete records and 17 columns, with Injury_Indicator as the target variable. Only 14 records are labeled as injuries, so the injury rate is 7%. The dataset has no missing values, which made the baseline cleaning process simple. Table I summarizes the dataset structure and the split used for the final experiments."
    )
    build_dataset_table(doc)
    add_body(
        doc,
        "The features cover several parts of the athlete profile. Age, height, and weight describe the athlete physically. Training_Intensity, Training_Hours_Per_Week, and Match_Count_Per_Week capture workload. Recovery_Days_Per_Week and Rest_Between_Events_Days capture recovery opportunity. Fatigue_Score, Performance_Score, Team_Contribution_Score, Load_Balance_Score, and ACL_Risk_Score provide higher-level performance and risk signals. Gender and Position are categorical variables. Athlete_ID is a pure identifier and does not carry predictive meaning."
    )
    add_body(
        doc,
        "To keep the baseline close to the SVM workflow shown in the presentation, Athlete_ID, Gender, and Position were removed from the numeric baseline frame. The all-feature extension kept Gender and Position and encoded them later in the preprocessing pipeline. The holdout split was stratified at 70% train, 15% validation, and 15% test, which produced 140 training records, 30 validation records, and 30 test records. Because of stratification, each of the validation and test splits contained 28 non-injury cases and 2 injury cases, while the training split contained 130 non-injury cases and 10 injury cases."
    )
    add_body(
        doc,
        "All experiments were implemented in Python using pandas, NumPy, scikit-learn, and matplotlib. A fixed random state of 42 was used for splitting and repeated cross-validation so that the results could be reproduced exactly from the repository. The project code also writes the same summary tables and figures used in this report, which helps keep the paper, presentation, and codebase consistent."
    )

    add_heading(doc, "V. METHODOLOGY AND IMPLEMENTATION")
    add_body(
        doc,
        "The baseline pipeline follows the parent paper most closely. After the numeric baseline frame was created, all predictor columns were standardized with StandardScaler so the SVM would not be dominated by variables with larger raw ranges. Recursive feature elimination (RFE) was then applied with a linear SVM to keep the 10 most useful predictors before the final classifier was trained. The final baseline model used an RBF-kernel SVM with C = 1.0, gamma = scale, and probability output enabled [6], [7]. The 10 selected baseline features were Height_cm, Training_Intensity, Training_Hours_Per_Week, Recovery_Days_Per_Week, Rest_Between_Events_Days, Fatigue_Score, Performance_Score, Team_Contribution_Score, Load_Balance_Score, and ACL_Risk_Score."
    )
    add_body(
        doc,
        "Those selected features are summarized in Table II. The pattern itself is informative. The baseline model concentrated on workload, recovery, fatigue, and direct risk signals rather than spreading its choices randomly across the dataset. That does not prove causal importance, but it does show that the compact SVM pipeline is at least focusing on variables that coaches and sports scientists would already monitor in practice."
    )
    build_selected_features_table(doc)
    add_body(
        doc,
        "The first comparison model was a Random Forest with 200 trees and class_weight set to balanced [8]. I used this model for two reasons. First, tree-based models are strong baselines on structured tabular data and do not require feature scaling. Second, they provide feature-importance rankings that are easier to explain to a coach or trainer than the boundary of an SVM. Even if the Random Forest did not improve the final score, it could still add value by showing which signals were driving the decision."
    )
    add_body(
        doc,
        "The second comparison model was an all-feature logistic regression extension. In this pipeline, all columns except Athlete_ID were kept. Numeric columns went through median imputation and scaling, and categorical columns went through most-frequent imputation and one-hot encoding before classification. The final classifier used balanced logistic regression with max_iter = 5000 and C = 5.0. Although the current dataset had no missing values, the imputation steps were still kept in the pipeline so that the code path would remain stable if the dataset changed later."
    )
    add_body(
        doc,
        "To make the recall tradeoff visible, I also evaluated a lower decision threshold on top of the baseline SVM probabilities. The default positive threshold for many classifiers is 0.50. In the warning-rule experiment, I lowered that threshold to 0.085. This was not presented as a fully tuned production threshold. Instead, it was used as a simple way to ask a practical question: what happens if the model is allowed to warn earlier, even if that means triggering more false positives?"
    )
    add_body(
        doc,
        "The evaluation strategy used both a fixed holdout split and repeated stratified cross-validation. The holdout split was kept because it matches the story shown in the project presentation and gives concrete examples such as the test confusion matrix. However, a 30-record test split can be noisy when only 2 injury cases are present. To avoid over-reading one small split, the project also ran repeated 5x10 stratified cross-validation across the full dataset. This gave a more stable view of how each model behaved across multiple train-test partitions."
    )
    add_body(
        doc,
        "Several evaluation metrics were used, but they do not all mean the same thing. Accuracy is the share of all predictions that were correct. Precision measures how many predicted injury cases were actually injuries. Recall measures how many real injury cases the model successfully found. F1-score balances precision and recall in a single number. Balanced accuracy averages recall across both classes, which makes it more informative than plain accuracy when the class counts are uneven. Finally, the confusion matrix shows the raw counts of true positives, false positives, true negatives, and false negatives. For injury prediction, recall and false negatives matter the most because a false negative means the model marked an at-risk athlete as safe [9]."
    )
    add_body(
        doc,
        "Implementation strategy was treated as part of the research design, not as a cleanup step after modeling. All final experiments were moved out of exploratory notebook cells and into a reusable Python module so the same split logic, feature handling, and metric calculations would be applied every time the code ran. The repository now contains a single experiment entry point, a lightweight runner script, saved result files, and unit tests. That structure matters because model comparison becomes unreliable if each candidate is being trained and evaluated through slightly different code paths."
    )
    add_body(
        doc,
        "The saved outputs were also chosen deliberately. The experiment runner exports a holdout leaderboard, a repeated cross-validation summary, a JSON summary of the full holdout experiment, a selected-features file for the baseline SVM, and figure files for the confusion matrix and Random Forest feature importance. Those artifacts are not just convenient for writing the report. They also make it easier to check whether later changes to the code accidentally alter the reported results. This was especially important in the final stage of the project because the report, slides, and repository all needed to tell the same story."
    )
    add_body(
        doc,
        "A final benefit of this implementation approach is that it clarifies what was exploratory and what was part of the verified pipeline. The baseline SVM, Random Forest, and logistic extension are all fully encoded in the experiment module and can be rerun directly. The lower-threshold warning rule is also saved as code, but it is described separately in the report because it is a decision rule layered on top of the SVM probabilities rather than a new model family. Keeping those distinctions explicit helps the paper stay honest about what was tested, what was tuned, and what still belongs to future work."
    )

    add_heading(doc, "VI. RESULTS")
    add_body(
        doc,
        "The validation split already hinted at the main issue. Both the baseline SVM and the Random Forest reached 0.9333 validation accuracy, but both predicted every validation case as non-injury, which means injury recall was 0.00 on that split. The all-feature logistic extension behaved differently. It reached 0.9667 validation accuracy, detected both validation injury cases, and produced one false positive. That result suggested early on that the project was not simply about finding the model with the biggest accuracy value."
    )
    add_body(
        doc,
        "That validation behavior is important because it reduces the chance that the test-set interpretation is just luck. If the logistic model had only looked better on one isolated holdout split, the result would have been much weaker. Instead, the validation split, the holdout split, and the cross-validation summary all point in the same direction: the safety-oriented models are more willing to capture the rare injury class, even when that choice costs some precision or raw accuracy."
    )
    add_body(
        doc,
        "The fixed holdout results are summarized in Table III. On the test split, the baseline SVM and Random Forest tied for the highest accuracy at 0.9667. At first glance, that looks like the strongest result in the report. The problem is that both models still missed one of the two injury cases, so their recall for the injury class remained 0.50. The logistic extension lowered accuracy to 0.9333, but it detected both injury cases and therefore reached recall of 1.00. The threshold rule did the same, although with a larger precision penalty."
    )
    build_holdout_table(doc, holdout_df)
    add_body(
        doc,
        "Figure 1 makes the holdout tradeoff easier to see. The accuracy leaders are not the recall leaders. The SVM and Random Forest sit highest on raw accuracy, but the logistic extension and the threshold rule are the only approaches that recover both injury cases on this split. Balanced accuracy also moves closer to the safety-oriented picture because it is less impressed by a model that does well only on the majority class."
    )
    add_figure(
        doc,
        tradeoff_chart,
        3.15,
        "Fig. 1. Holdout comparison of accuracy, balanced accuracy, and injury recall. The chart shows why the best accuracy score is not automatically the safest model."
    )
    add_body(
        doc,
        "The baseline SVM confusion matrix in Fig. 2 shows the exact point where the interpretation changes. The model correctly identified all 28 non-injury cases and one injury case, but it missed the other injury case. On a larger test set, one false negative might look like a minor detail. On this test split, it cuts injury recall in half. That single miss is enough to change the model from a strong classifier on paper to a questionable early-warning tool in practice."
    )
    add_figure(
        doc,
        confusion_chart,
        3.0,
        "Fig. 2. Baseline SVM confusion matrix on the held-out test split. One false negative is enough to reduce injury recall from 1.00 to 0.50."
    )
    add_body(
        doc,
        "The Random Forest did not improve the holdout recall, but it was still useful because it exposed which variables mattered most. Figure 3 shows that ACL_Risk_Score was the strongest feature by a clear margin, followed by Load_Balance_Score, Fatigue_Score, and Recovery_Days_Per_Week. This ranking is plausible from a sports analytics perspective because those variables are directly connected to strain, recovery quality, and injury exposure. In other words, the model was not basing its decisions on obviously meaningless signals."
    )
    add_figure(
        doc,
        importance_chart,
        3.15,
        "Fig. 3. Random Forest feature importance from the numeric baseline frame. ACL risk, load balance, fatigue, and recovery were the strongest signals."
    )
    add_body(
        doc,
        "The repeated cross-validation results in Table IV provide the broader picture. Across 50 stratified folds, Random Forest had the highest mean accuracy at 0.9560, which confirms that the strong holdout score was not a complete fluke. However, its mean recall was still only 0.4000. The all-feature logistic extension had slightly lower mean accuracy at 0.9435, but it achieved mean balanced accuracy of 0.8588 and mean recall of 0.7600. A compact logistic model with RFE8 pushed recall even higher to 0.8633, but gave up more accuracy. The baseline SVM had the weakest mean recall by far at 0.1633. Taken together, these cross-validation results support the same conclusion as the holdout analysis: accuracy alone overstates the value of the baseline SVM in this injury-prediction setting."
    )
    build_cv_table(doc, cv_df)

    add_heading(doc, "VII. DISCUSSION")
    add_body(
        doc,
        "The answer to the research question is mixed. Yes, an SVM-based pipeline can produce strong top-line scores on this dataset, and the parent-paper style workflow is reproducible. But once missed injuries are treated as a central outcome, the plain baseline is much less convincing. The SVM and Random Forest both look excellent on accuracy and still fail to recover all injury cases on the fixed holdout split. That is not a small technicality. It is the core practical weakness of the baseline framing."
    )
    add_body(
        doc,
        "This finding does not mean the parent paper was wrong to use SVM. Instead, it shows how much the interpretation depends on the evaluation goal and the dataset structure. Li [1] used a larger big-data framing and reported strong performance. My dataset is smaller, highly imbalanced, and limited to 200 records. Under those conditions, a model that mostly predicts non-injury can still look impressive unless the analysis directly checks recall and false negatives. That observation is consistent with the broader literature, which repeatedly points to imbalance, interpretation, and contextual validity as major concerns in sports injury modeling [2]-[4], [9]."
    )
    add_body(
        doc,
        "The comparison models also help clarify what kind of system might be more useful in practice. Random Forest matched the strongest holdout accuracy and gave a cleaner explanation of feature importance, which makes it attractive when interpretability matters. The logistic extension, on the other hand, was the strongest safety-oriented model in this project because it repeatedly recovered more injury cases. The threshold rule is the simplest version of the same idea: accept more false alarms in exchange for fewer missed high-risk athletes. A coaching staff may not want constant noise, but it may still prefer a noisier warning system over one that quietly misses an athlete who should have been flagged."
    )
    add_body(
        doc,
        "That tradeoff also suggests a practical division between model roles. A staff member making weekly planning decisions might prefer a more conservative and interpretable model such as Random Forest, especially if the output needs to be explained quickly. A sports scientist or athletic trainer screening for early warning signs might prefer the logistic extension or a lower-threshold SVM because those approaches are more willing to surface possible risk. The same dataset can therefore support different decisions depending on whether the goal is explanation, triage, or aggressive prevention."
    )
    add_body(
        doc,
        "One useful outcome of the project is that it shifts the work away from model shopping and toward better framing. The biggest lesson was not that one algorithm permanently beats another. The bigger lesson was that injury prediction needs the right success definition. If the true goal is early intervention, then the model should be judged on how often it catches risk in time, not only on how often it agrees with the majority class."
    )

    add_heading(doc, "VIII. LIMITATIONS AND FUTURE WORK")
    add_body(
        doc,
        "This study has clear limitations. The dataset is small, and only 14 of the 200 records belong to the injury class. That means every positive example carries a lot of weight, and a single classification error can noticeably move the reported metrics. The current dataset is also a classroom dataset rather than a multi-season professional tracking dataset, so the results should not be overgeneralized to every sport or competitive level."
    )
    add_body(
        doc,
        "A second limitation is that the current model views each athlete record as a mostly static snapshot. That is a major simplification. In real training environments, injury risk usually builds across time through workload accumulation, fatigue, recovery quality, and repeated exposure. A stronger next version of this project should create rolling 7-day and 28-day workload features, compare acute and chronic load patterns, and test true sequential models such as LSTM, GRU, or other time-series approaches when the data supports them [5]."
    )
    add_body(
        doc,
        "Future work should also improve the evaluation setup. The threshold rule in this report was intentionally simple and should not be treated as fully optimized. A stronger study would tune thresholds on the validation split, evaluate probability calibration, inspect precision-recall curves directly, and test the system on an external dataset. For practical use, future models should also include coach-friendly explanation tools so that high risk warnings can be connected to workload, fatigue, recovery, and known biomechanical factors."
    )
    add_body(
        doc,
        "There is also room to improve how the positive class is handled during model training. In this report, the main imbalance responses were class weighting and a recall-first threshold check. Future experiments could compare those choices against resampling methods, cost-sensitive learning, or probability calibration strategies that are better suited to rare-event classification. That would make the next version of the study stronger both methodologically and practically."
    )

    add_heading(doc, "IX. CONCLUSION")
    add_body(
        doc,
        "This project started from an SVM-based parent paper and turned that idea into a reproducible student-built pipeline for athlete injury prediction. On the fixed holdout split, the baseline SVM and Random Forest both reached 96.67% accuracy, but both also missed one of the two injury cases. That gap matters because injury prediction is not only about being correct most of the time. It is about missing as few risky athletes as possible."
    )
    add_body(
        doc,
        "The strongest final takeaway is that a safer evaluation lens changes the whole story. When recall, balanced accuracy, and false negatives were treated seriously, the all-feature logistic extension and the lower-threshold warning rule became much more valuable than the raw accuracy ranking suggested. The project therefore supports machine learning as a useful tool for injury analysis, but it also shows that future systems should be time-aware, recall-conscious, and explainable enough for real sports decision-making."
    )

    add_heading(doc, "REFERENCES")
    references = [
        "[1] W. Li, \"A Big Data Approach to Forecast Injuries in Professional Sports Using Support Vector Machine,\" Mobile Networks and Applications, 2024, doi: 10.1007/s11036-024-02377-x.",
        "[2] H. Van Eetvelde et al., \"Machine learning methods in sport injury prediction and prevention: a systematic review,\" Journal of Experimental Orthopaedics, vol. 8, art. no. 27, 2021, doi: 10.1186/s40634-021-00346-x.",
        "[3] A. Majumdar et al., \"Machine Learning for Understanding and Predicting Injuries in Football,\" Sports Medicine - Open, vol. 8, art. no. 73, 2022, doi: 10.1186/s40798-022-00465-4.",
        "[4] A. Pareek, D. H. Ro, J. Karlsson, and R. K. Martin, \"Machine learning/artificial intelligence in sports medicine: state of the art and future directions,\" Journal of ISAKOS, vol. 9, no. 4, pp. 635-644, 2024, doi: 10.1016/j.jisako.2024.01.013.",
        "[5] X. Ye, Y. Huang, Z. Bai, and Y. Wang, \"A novel approach for sports injury risk prediction: based on time-series image encoding and deep learning,\" Frontiers in Physiology, vol. 14, art. no. 1174525, 2023, doi: 10.3389/fphys.2023.1174525.",
        "[6] C. Cortes and V. Vapnik, \"Support-vector networks,\" Machine Learning, vol. 20, no. 3, pp. 273-297, 1995.",
        "[7] I. Guyon, J. Weston, S. Barnhill, and V. Vapnik, \"Gene selection for cancer classification using support vector machines,\" Machine Learning, vol. 46, nos. 1-3, pp. 389-422, 2002.",
        "[8] L. Breiman, \"Random forests,\" Machine Learning, vol. 45, no. 1, pp. 5-32, 2001.",
        "[9] T. Saito and M. Rehmsmeier, \"The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets,\" PLoS ONE, vol. 10, no. 3, art. no. e0118432, 2015, doi: 10.1371/journal.pone.0118432.",
    ]
    for reference in references:
        add_reference(doc, reference)

    doc.save(OUT_DOCX)
    shutil.copy2(OUT_DOCX, OUTPUT_COPY)
    shutil.copy2(OUT_DOCX, DOWNLOAD_COPY)
    return {
        "report_docx": str(OUT_DOCX),
        "output_copy": str(OUTPUT_COPY),
        "download_copy": str(DOWNLOAD_COPY),
        "tradeoff_chart": str(tradeoff_chart),
    }


if __name__ == "__main__":
    print(json.dumps(build_report(), indent=2))
